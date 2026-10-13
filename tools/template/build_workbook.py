# Builds one brand's content-calendar workbook for a financial year (April to March).
# Every brand uses exactly this structure; only the data differs. This file is the single
# definition of the workbook: the website's import must read exactly what it writes.
#
# Run from the tools folder with its virtual environment active (see tools/README.md):
#   python template/build_workbook.py --brand "Emkay Global" --fy 2027 --bau-start 2026-11
#   python template/build_workbook.py --from "Emkay IM - Content calendar FY27.xlsx" --fy 2028
#
# --from carries the brand, purpose, BAU start, comms tasks, framework and lists over from
# last year's workbook, so next year's file starts with the options already in place.

import calendar, datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

TEMPLATE_VERSION = '1.0'
FONT = 'Arial'
MAX_ROW = 400  # validation and highlighting reach this far, so new rows are covered

INK = '1F2937'
HEAD_FILL = PatternFill('solid', fgColor='1F3A5F')
AUTO_FILL = PatternFill('solid', fgColor='EEF0F3')
FLAG_FILL = PatternFill('solid', fgColor='FFF2A8')
WEEKEND_FILL = PatternFill('solid', fgColor='F7F7F2')
RED_FILL = PatternFill('solid', fgColor='F8C9C4')
AMBER_FILL = PatternFill('solid', fgColor='FCE1B8')
THIN = Side(style='thin', color='D0D5DD')

STATUSES = [
    ('Idea', 'A possible post. Not confirmed yet.'),
    ('Planned', 'Confirmed for this date. Work has not started.'),
    ('Copy in progress', 'The copy is being written.'),
    ('Design in progress', 'The brief is with design.'),
    ('Ready to publish', 'Copy and design are approved.'),
    ('Published', 'Live. Paste the post link in each platform column.'),
    ('Dropped', 'Will not be published. Keep the row so the history is complete.'),
]
FORMATS = ['Carousel', 'Static', 'Infographic', 'Video', 'Text', 'Article']
PLATFORMS = ['LinkedIn', 'Instagram', 'YouTube', 'X', 'WhatsApp']

# Month sheet columns: (header, width, kind). kind: input | auto | link
MONTH_COLUMNS = [
    ('Date', 12, 'input'),
    ('Day (auto)', 8, 'auto'),
    ('Comms task (auto)', 11, 'auto'),
    ('Pillar (auto)', 22, 'auto'),
    ('Peg', 30, 'input'),
    ('Series', 18, 'input'),
    ('Title', 40, 'input'),
    ('Details', 40, 'input'),
    ('Format', 12, 'input'),
    ('Owner', 16, 'input'),
    ('Status', 17, 'input'),
    *[(p, 13, 'link') for p in PLATFORMS],
    ('Copy link', 13, 'link'),
    ('Long-form', 13, 'link'),
    ('Reference', 13, 'link'),
    ('Moment marketing', 24, 'input'),
]
COL = {h: i + 1 for i, (h, _, _) in enumerate(MONTH_COLUMNS)}
KEYS = {  # short keys used by the migration
    'date': 'Date', 'peg': 'Peg', 'series': 'Series', 'title': 'Title', 'details': 'Details',
    'format': 'Format', 'owner': 'Owner', 'status': 'Status', 'copy': 'Copy link',
    'longform': 'Long-form', 'reference': 'Reference', 'moment': 'Moment marketing',
    **{p: p for p in PLATFORMS},
}


def letter(header):
    from openpyxl.utils import get_column_letter
    return get_column_letter(COL[header])


def fy_months(start_year):
    return [(start_year + (m < 4), m) for m in list(range(4, 13)) + [1, 2, 3]]


def sheet_name(y, m):
    return f'{calendar.month_abbr[m]} {y}'


def table_name(y, m):
    return f'{calendar.month_abbr[m]}_{y}'


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), color=kw.pop('color', INK), **kw)


def style_header(ws, row, first, last, fill=HEAD_FILL):
    for c in range(first, last + 1):
        cell = ws.cell(row, c)
        cell.font = font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30


def add_table(ws, name, ref, formulas=None):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    t._initialise_columns()
    from openpyxl.utils.cell import range_boundaries
    c1, r1, _, _ = range_boundaries(ref)
    for i, col in enumerate(t.tableColumns):
        col.name = str(ws.cell(r1, c1 + i).value)
    for idx, f in (formulas or {}).items():
        t.tableColumns[idx].calculatedColumnFormula = TableFormula(attr_text=f)
    ws.add_table(t)
    return t


def list_name(wb, name, sheet, col):
    ref = f"OFFSET('{sheet}'!${col}$2,0,0,MAX(1,COUNTA('{sheet}'!${col}:${col})-1),1)"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def list_dv(ws, rng, source, title, msg):
    dv = DataValidation(type='list', formula1=source, allow_blank=True, showErrorMessage=True,
                        errorTitle=title, error=msg, showInputMessage=False)
    dv.add(rng)
    ws.add_data_validation(dv)


def flag(cell, text):
    cell.fill = FLAG_FILL
    note = Comment(text, 'Migration review')
    note.width, note.height = 320, 140
    cell.comment = note


def build(path, brand, fy_start, bau_start=None, purpose='', tasks=(), pegs=(), team=(), series=(), rows=(),
          flags=(), peg_map=(), formats=FORMATS, default_poc='BAU lead'):
    """tasks: (code, name, counts_as_bau, target or None, active, notes); pegs: (code, pillar, peg, active, notes, link[, flag]);
    series: (name, description); rows: dicts keyed by KEYS with 'date' a datetime.date and
    optional '_flags': {key: text}; flags: general (where, issue, action) rows for the review sheet."""
    wb = Workbook()
    fy = f'FY{(fy_start + 1) % 100:02d}'
    months = fy_months(fy_start)
    review = []  # (sheet, cell, title, issue, action)

    # ---------- Start here ----------
    ws = wb.active
    ws.title = 'Start here'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 95
    ws['A1'] = f'{brand} content calendar'
    ws['A1'].font = font(size=16, bold=True)
    info = [('Brand', brand), ('Financial year', f'{fy} (April {fy_start} to March {fy_start + 1})'),
            ('Template version', TEMPLATE_VERSION), ('Purpose', purpose or ''), ('BAU started', bau_start)]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(i, 1, k).font = font(bold=True)
        c = ws.cell(i, 2, v)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws['B6'].protection = Protection(locked=False)
    ws['B7'].number_format = 'mmmm yyyy'
    ws['B7'].alignment = Alignment(horizontal='left')
    ws['B7'].comment = Comment('The month this brand started using the BAU framework. The Summary ignores earlier months; '
                               'posts from before stay in the calendar as a baseline for analytics.', 'Template')
    wb.defined_names['BAUStart'] = DefinedName('BAUStart', attr_text="'Start here'!$B$7")
    guide = [
        ('How to fill it in', None),
        ('1', 'Open the month sheet. Every day already has a row. To add a second post on a day, right-click the row → Insert → Table Rows Above, then type the date.'),
        ('2', 'Pick the Peg from the dropdown. Comms task and Pillar fill themselves in from the Framework sheet. Never type in grey columns.'),
        ('3', 'Pick Series (if any), Format, Owner and Status from their dropdowns. Typing something that is not in the list is refused.'),
        ('4', 'In each platform column type Planned, then replace it with the post URL once the post is live. WhatsApp also accepts Sent.'),
        ('5', 'A row counts as a post once it has a Title or a Peg. A row with only Moment marketing is a reminder, not a post.'),
        ('', 'Copy link, Long-form and Reference only accept links (starting with http).'),
        ('6', 'Never rename, add or delete month sheets, and keep summaries off the month sheets: the Summary sheet does that automatically.'),
        ('Changing the options', None),
        ('Framework', 'Add a peg as a new row. To stop using a peg, set Active to No instead of deleting it: it leaves the dropdown, but older posts still count towards it.'),
        ('Retire a pillar', 'Set Active to No on every peg in that pillar.'),
        ('Retire a comms task', 'Set Active to No on the Comms tasks sheet. All its pegs leave the dropdown; past posts and the Summary keep their history.'),
        ('Mid-year changes', 'Do not move an existing peg to another pillar or comms task: that re-files every past post. Retire it and add a new peg instead. Restructure at the start of a financial year.'),
        ('Comms tasks', 'Codes, whether they count as BAU, and the target share of BAU posts. Compliance posts use a comms task with Counts as BAU = No, so they are tracked but kept out of the BAU mix. BAU targets should add up to 100%.'),
        ('Lists', 'Team members, series (branded properties), formats and statuses. Add a row to add an option.'),
        ('Colour key', None),
        ('White', 'You type or pick a value.'),
        ('Grey', 'Automatic. Do not type here.'),
        ('Yellow', 'Needs your review. Hover for the note.'),
        ('Orange', 'Something is missing: for example a post with no peg, or a published post with no link.'),
        ('Red', 'Something is broken: a peg that is not in the framework, or a formula that was typed over.'),
    ]
    r = 9
    for k, v in guide:
        if v is None:
            r += 1
            ws.cell(r, 1, k).font = font(size=12, bold=True)
        else:
            ws.cell(r, 1, k).font = font(bold=True)
            c = ws.cell(r, 2, v)
            c.font = font()
            c.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    for label, fill in [('White', None), ('Grey', AUTO_FILL), ('Yellow', FLAG_FILL), ('Orange', AMBER_FILL), ('Red', RED_FILL)]:
        for row in ws.iter_rows(min_row=9, max_row=r, max_col=1):
            if row[0].value == label and fill:
                row[0].fill = fill
    ws.protection.sheet = True

    # ---------- Month sheets ----------
    by_month = {}
    for row in rows:
        by_month.setdefault((row['date'].year, row['date'].month), []).append(row)
    hdrs = [h for h, _, _ in MONTH_COLUMNS]
    for (y, m) in months:
        ws = wb.create_sheet(sheet_name(y, m))
        tname = table_name(y, m)
        if bau_start and datetime.date(y, m, calendar.monthrange(y, m)[1]) < bau_start:
            ws.sheet_properties.tabColor = 'B0B7C3'
        for i, (h, w, kind) in enumerate(MONTH_COLUMNS, start=1):
            ws.cell(1, i, h)
            ws.column_dimensions[letter(h)].width = w
        style_header(ws, 1, 1, len(hdrs))
        for h in ('Day (auto)', 'Comms task (auto)', 'Pillar (auto)'):
            ws.cell(1, COL[h]).fill = PatternFill('solid', fgColor='5B6B80')
        ws.freeze_panes = 'C2'

        this = sorted(by_month.get((y, m), []), key=lambda x: x['date'])
        days = calendar.monthrange(y, m)[1]
        out = []
        for d in range(1, days + 1):
            date = datetime.date(y, m, d)
            todays = [x for x in this if x['date'] == date]
            out.extend(todays or [{'date': date}])
        f_day = f'IF({tname}[[#This Row],[Date]]="","",TEXT({tname}[[#This Row],[Date]],"ddd"))'
        f_task = (f'IF({tname}[[#This Row],[Peg]]="","",IFERROR(INDEX(Pegs[Comms task code],'
                  f'MATCH({tname}[[#This Row],[Peg]],Pegs[Peg],0)),"Not in framework"))')
        f_pillar = (f'IF({tname}[[#This Row],[Peg]]="","",IFERROR(INDEX(Pegs[Pillar],'
                    f'MATCH({tname}[[#This Row],[Peg]],Pegs[Peg],0)),"Not in framework"))')
        for r, item in enumerate(out, start=2):
            for key, h in KEYS.items():
                v = item.get(key)
                if v in (None, ''):
                    continue
                c = ws.cell(r, COL[h], v)
                if isinstance(v, str) and v.startswith('http'):
                    c.hyperlink = v
                    c.font = font(color='1D4ED8', underline='single')
            ws.cell(r, COL['Day (auto)'], '=' + f_day)
            ws.cell(r, COL['Comms task (auto)'], '=' + f_task)
            ws.cell(r, COL['Pillar (auto)'], '=' + f_pillar)
            ws.cell(r, 1).number_format = 'dd mmm yyyy'
            for key, text in item.get('_flags', {}).items():
                text, who = text if isinstance(text, tuple) else (text, item.get('owner') or default_poc)
                h = KEYS.get(key, key)
                cell = ws.cell(r, COL[h])
                flag(cell, f'For {who}: {text}')
                review.append((ws.title, cell.coordinate, item.get('title') or item.get('peg') or '', text, who))
        last = max(2, len(out) + 1)
        for row in ws.iter_rows(min_row=2, max_row=last):
            for c in row:
                if not c.font or c.font.name != FONT:
                    c.font = font(color=c.font.color.rgb if c.hyperlink else INK, underline='single' if c.hyperlink else None)
                c.alignment = Alignment(vertical='top', wrap_text=c.column in (COL['Title'], COL['Details'], COL['Moment marketing']))
                c.border = Border(bottom=THIN)
                if hdrs[c.column - 1].endswith('(auto)'):
                    c.fill = AUTO_FILL
        add_table(ws, tname, f'A1:{letter("Moment marketing")}{last}',
                  {COL['Day (auto)'] - 1: f_day, COL['Comms task (auto)'] - 1: f_task, COL['Pillar (auto)'] - 1: f_pillar})

        # Validation
        first, end = datetime.date(y, m, 1), datetime.date(y, m, days)
        dv = DataValidation(type='date', operator='between',
                            formula1=f'DATE({y},{m},1)', formula2=f'DATE({y},{m},{days})',
                            allow_blank=True, showErrorMessage=True, errorTitle='Date outside this month',
                            error=f'This sheet only takes dates from {first:%d %b %Y} to {end:%d %b %Y}, typed as a date (for example {first:%d/%m/%Y}).')
        dv.add(f'A2:A{MAX_ROW}')
        ws.add_data_validation(dv)
        for h, src, what in [('Peg', 'PegList', 'peg from the Framework sheet'), ('Series', 'SeriesList', 'series from the Lists sheet'),
                             ('Format', 'FormatList', 'format from the Lists sheet'), ('Owner', 'TeamList', 'team member from the Lists sheet'),
                             ('Status', 'StatusList', 'status from the Lists sheet')]:
            col = letter(h)
            list_dv(ws, f'{col}2:{col}{MAX_ROW}', src, f'Not a valid {h.lower()}', f'Pick a {what}. To add a new option, add it there first.')
        for h in ('Copy link', 'Long-form', 'Reference'):
            col = letter(h)
            dv = DataValidation(type='custom', formula1=f'LEFT({col}2,4)="http"', allow_blank=True, showErrorMessage=True,
                                errorTitle='Links only', error='Paste the full link (starting with http).')
            dv.add(f'{col}2:{col}{MAX_ROW}')
            ws.add_data_validation(dv)
        for p in PLATFORMS:
            col = letter(p)
            extra = f',{col}2="Sent"' if p == 'WhatsApp' else ''
            dv = DataValidation(type='custom', formula1=f'OR({col}2="Planned"{extra},LEFT({col}2,4)="http")', allow_blank=True,
                                showErrorMessage=True, errorTitle='Planned or a link',
                                error='Type Planned, or paste the full post link (starting with http).' + (' WhatsApp also accepts Sent.' if extra else ''))
            dv.add(f'{col}2:{col}{MAX_ROW}')
            ws.add_data_validation(dv)

        # Highlighting
        A, E, G, K = 'A', letter('Peg'), letter('Title'), letter('Status')
        L, P = letter('LinkedIn'), letter('WhatsApp')
        is_post = f'OR(${G}2<>"",${E}2<>"")'
        for h in ('Day (auto)', 'Comms task (auto)', 'Pillar (auto)'):
            col = letter(h)
            ws.conditional_formatting.add(f'{col}2:{col}{MAX_ROW}', FormulaRule(
                formula=[f'AND(COUNTA($A2,${E}2,${G}2)>0,NOT(_xlfn.ISFORMULA({col}2)))'], fill=RED_FILL, stopIfTrue=True))
        ws.conditional_formatting.add(f'C2:D{MAX_ROW}', FormulaRule(formula=['C2="Not in framework"'], fill=RED_FILL, font=Font(name=FONT, bold=True, color='9B1C1C'), stopIfTrue=True))
        ws.conditional_formatting.add(f'{E}2:{E}{MAX_ROW}', FormulaRule(formula=[f'AND(${G}2<>"",${E}2="",${K}2<>"Dropped")'], fill=AMBER_FILL))
        ws.conditional_formatting.add(f'{K}2:{K}{MAX_ROW}', FormulaRule(formula=[f'AND({is_post},${K}2="")'], fill=AMBER_FILL))
        ws.conditional_formatting.add(f'{K}2:{K}{MAX_ROW}', FormulaRule(
            formula=[f'AND(${K}2="Published",COUNTIF(${L}2:${P}2,"http*")=0)'], fill=AMBER_FILL))
        ws.conditional_formatting.add(f'A2:{letter("Moment marketing")}{MAX_ROW}', FormulaRule(
            formula=['AND($A2<>"",WEEKDAY($A2,2)>5)'], fill=WEEKEND_FILL))

    # ---------- Summary ----------
    ws = wb.create_sheet('Summary')
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'{brand}: posts per comms task, {fy}'
    ws['A1'].font = font(size=14, bold=True)
    ws['A2'] = ('Automatic. Counts every row with a Peg, except Dropped. Share is the share of BAU posts, compared with the target '
                'on the Comms tasks sheet. Comms tasks that do not count as BAU (such as Compliance) are counted but left out of the share. '
                'Months before the BAU start date (Start here sheet) are left blank.')
    ws['A2'].font = font(italic=True, color='667085')
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A2:P2')
    ws.row_dimensions[2].height = 42
    n_tasks = max(8, len(tasks) + 3)
    month_labels = [sheet_name(y, m) for y, m in months]
    M0 = 5  # first month column (E)
    FYC = M0 + len(months)

    def pre(y, m):
        return f'DATE({y},{m},1)<DATE(YEAR(BAUStart),MONTH(BAUStart),1)'

    def heads(r):
        for i, t in enumerate(['Code', 'Comms task', 'BAU?', 'Target'] + month_labels + ['Full year'], start=1):
            ws.cell(r, i, t)
        style_header(ws, r, 1, FYC)

    def task_cols(r, k):
        idx = k + 1
        ws.cell(r, 1, f'=IFERROR(INDEX(CommsTasks[Code],{idx}),"")')
        ws.cell(r, 2, f'=IF($A{r}="","",INDEX(CommsTasks[Comms task],{idx}))')
        ws.cell(r, 3, f'=IF($A{r}="","",IF(INDEX(CommsTasks[Counts as BAU],{idx})="No","No","Yes"))')
        ws.cell(r, 4, f'=IF(OR($A{r}="",$C{r}="No"),"",IF(INDEX(CommsTasks[Target share],{idx})="","",INDEX(CommsTasks[Target share],{idx})))')
        ws.cell(r, 4).number_format = '0%'

    col = lambda j: ws.cell(1, j).column_letter
    # Block 1: counts
    ws.cell(4, 1, 'Number of posts').font = font(size=12, bold=True)
    heads(5)
    c_first, c_last = 6, 5 + n_tasks
    for k in range(n_tasks):
        r = c_first + k
        task_cols(r, k)
        for j, (y, m) in enumerate(months):
            t = table_name(y, m)
            ws.cell(r, M0 + j, f'=IF(OR($A{r}="",{pre(y, m)}),"",COUNTIFS({t}[Comms task (auto)],$A{r},{t}[Status],"<>Dropped"))')
        ws.cell(r, FYC, f'=IF($A{r}="","",SUM({col(M0)}{r}:{col(FYC - 1)}{r}))').font = font(bold=True)
    bau_r, non_r = c_last + 1, c_last + 2
    ws.cell(bau_r, 2, 'BAU posts').font = font(bold=True)
    ws.cell(non_r, 2, 'Not BAU (e.g. compliance)').font = font(bold=True)
    for j in range(len(months) + 1):
        L = col(M0 + j)
        ws.cell(bau_r, M0 + j, f'=SUMIF($C${c_first}:$C${c_last},"Yes",{L}{c_first}:{L}{c_last})').font = font(bold=True)
        ws.cell(non_r, M0 + j, f'=SUMIF($C${c_first}:$C${c_last},"No",{L}{c_first}:{L}{c_last})').font = font(bold=True)
    # Block 2: share of BAU posts
    s_top = non_r + 2
    ws.cell(s_top, 1, 'Share of BAU posts (compare with Target)').font = font(size=12, bold=True)
    heads(s_top + 1)
    s_first = s_top + 2
    s_last = s_first + n_tasks - 1
    for k in range(n_tasks):
        r, cr = s_first + k, c_first + k
        task_cols(r, k)
        for j in range(len(months) + 1):
            L = col(M0 + j)
            c = ws.cell(r, M0 + j, f'=IF(OR($A{r}="",$C{r}="No",{L}{cr}="",N({L}${bau_r})=0),"",{L}{cr}/{L}${bau_r})')
            c.number_format = '0%'
            if j == len(months):
                c.font = font(bold=True)
    ws.conditional_formatting.add(f'{col(M0)}{s_first}:{col(FYC)}{s_last}', FormulaRule(
        formula=[f'AND(ISNUMBER({col(M0)}{s_first}),ISNUMBER($D{s_first}),ABS({col(M0)}{s_first}-$D{s_first})>0.1)'], fill=AMBER_FILL))
    # Block 3: data checks
    q_top = s_last + 3
    ws.cell(q_top, 1, 'Data checks (should all be 0)').font = font(size=12, bold=True)
    ws.cell(q_top + 1, 1, 'Check')
    for j, lab in enumerate(month_labels + ['Full year']):
        ws.cell(q_top + 1, M0 + j, lab)
    style_header(ws, q_top + 1, 1, FYC)
    checks = [
        ('Posts with no peg', lambda t: f'COUNTIFS({t}[Title],"<>",{t}[Peg],"",{t}[Status],"<>Dropped")'),
        ('Pegs not in the framework', lambda t: f'COUNTIF({t}[Comms task (auto)],"Not in framework")'),
        ('Posts with no status', lambda t: f'COUNTIFS({t}[Title],"<>",{t}[Status],"")+COUNTIFS({t}[Title],"",{t}[Peg],"<>",{t}[Status],"")'),
        ('Posts with no owner', lambda t: f'COUNTIFS({t}[Title],"<>",{t}[Owner],"")+COUNTIFS({t}[Title],"",{t}[Peg],"<>",{t}[Owner],"")'),
    ]
    for i, (lab, f) in enumerate(checks):
        r = q_top + 2 + i
        ws.cell(r, 1, lab)
        for j, (y, m) in enumerate(months):
            ws.cell(r, M0 + j, f'=IF({pre(y, m)},"",{f(table_name(y, m))})')
        ws.cell(r, FYC, f'=SUM({col(M0)}{r}:{col(FYC - 1)}{r})').font = font(bold=True)
    ws.conditional_formatting.add(f'{col(M0)}{q_top + 2}:{col(FYC)}{q_top + 1 + len(checks)}',
                                  FormulaRule(formula=[f'{col(M0)}{q_top + 2}>0'], fill=AMBER_FILL))
    for row in ws.iter_rows(min_row=4, max_row=q_top + 1 + len(checks)):
        for c in row:
            if c.font.name != FONT:
                c.font = font(bold=c.font.b)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 7
    ws.column_dimensions['D'].width = 8
    for j in range(len(months) + 1):
        ws.column_dimensions[col(M0 + j)].width = 10
    ws.freeze_panes = 'E1'
    ws.protection.sheet = True

    # ---------- Comms tasks ----------
    ws = wb.create_sheet('Comms tasks')
    heads_ = [('Code', 10), ('Comms task', 28), ('Counts as BAU', 13), ('Target share', 13), ('Active', 9), ('Notes', 60)]
    for i, (h, w) in enumerate(heads_, start=1):
        ws.cell(1, i, h)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    style_header(ws, 1, 1, len(heads_))
    for r, (code, name, bau, target, active, notes) in enumerate(tasks, start=2):
        for c, v in enumerate((code, name, bau, target, active, notes), start=1):
            if v not in (None, ''):
                ws.cell(r, c, v)
    last = max(2, len(tasks) + 1)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = font()
            c.alignment = Alignment(vertical='top', wrap_text=True)
        row[3].number_format = '0%'
    add_table(ws, 'CommsTasks', f'A1:F{last}')
    for colx in ('C', 'E'):
        dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True, showErrorMessage=True)
        dv.add(f'{colx}2:{colx}{MAX_ROW}')
        ws.add_data_validation(dv)
    r = last + 2
    ws.cell(r, 2, 'BAU targets add up to').font = font(bold=True)
    ws.cell(r, 4, '=SUMIFS(CommsTasks[Target share],CommsTasks[Counts as BAU],"<>No")').number_format = '0%'
    ws.cell(r, 4).font = font(bold=True)
    ws.conditional_formatting.add(f'D{r}', FormulaRule(formula=[f'AND(D{r}>0,ROUND(D{r},4)<>1)'], fill=RED_FILL))
    ws.conditional_formatting.add(f'A2:A{MAX_ROW}', FormulaRule(formula=['AND(A2<>"",COUNTIF($A$2:$A$100,A2)>1)'], fill=RED_FILL))
    ws.conditional_formatting.add(f'A2:F{MAX_ROW}', FormulaRule(formula=['$E2="No"'], font=Font(name=FONT, color='98A2B3', strike=True)))
    ws.cell(r + 2, 1, 'Retire a comms task by setting Active to No. Do not delete it: past posts still count towards it.').font = font(italic=True, color='667085')
    ws.freeze_panes = 'A2'
    list_name(wb, 'TaskList', 'Comms tasks', 'A')

    # ---------- Framework ----------
    ws = wb.create_sheet('Framework')
    heads = [('Comms task code', 12), ('Pillar', 28), ('Peg', 36), ('Active', 9), ('Notes', 55), ('Example', 14), ('In dropdown (auto)', 11)]
    for i, (h, w) in enumerate(heads, start=1):
        ws.cell(1, i, h)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    style_header(ws, 1, 1, len(heads))
    for r, (code, pillar, peg, active, notes, link, *pf) in enumerate(pegs, start=2):
        for c, v in enumerate((code, pillar, peg, active, notes, link), start=1):
            if v:
                ws.cell(r, c, v)
        if link:
            ws.cell(r, 6).hyperlink = link
        if pf and pf[0]:
            flag(ws.cell(r, 3), f'For BAU lead: {pf[0]}')
            review.append(('Framework', f'C{r}', peg, pf[0], 'BAU lead'))
    last = max(2, len(pegs) + 1)
    # Running number for pegs that are active and whose comms task is active; feeds the Peg dropdown.
    f_dd = ('IF(AND(C2<>"",D2<>"No",IFERROR(INDEX(CommsTasks[Active],MATCH(A2,CommsTasks[Code],0))<>"No",TRUE)),'
            'MAX(G$1:G1)+1,"")')
    for r in range(2, last + 1):
        ws.cell(r, 7, '=' + f_dd.replace('C2', f'C{r}').replace('D2', f'D{r}').replace('A2', f'A{r}').replace('G1)', f'G{r - 1})'))
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = font(color='1D4ED8', underline='single') if c.hyperlink else font()
            c.alignment = Alignment(vertical='top', wrap_text=True)
        row[6].fill = AUTO_FILL
    ws.cell(1, 7).fill = PatternFill('solid', fgColor='5B6B80')
    add_table(ws, 'Pegs', f'A1:G{last}', {6: f_dd})
    list_dv(ws, f'A2:A{MAX_ROW}', 'TaskList', 'Unknown comms task', 'Use a code from the Comms tasks sheet.')
    dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True, showErrorMessage=True)
    dv.add(f'D2:D{MAX_ROW}')
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(f'C2:C{MAX_ROW}', FormulaRule(formula=['AND(C2<>"",COUNTIF($C$2:$C$400,C2)>1)'], fill=RED_FILL))
    ws.conditional_formatting.add(f'A2:F{MAX_ROW}', FormulaRule(formula=['$D2="No"'], font=Font(name=FONT, color='98A2B3', strike=True)))
    ws.freeze_panes = 'A2'
    ws.cell(last + 2, 1, 'Each peg appears once. Pick pegs on the month sheets; Comms task and Pillar follow from this sheet. '
                         'Retire a peg by setting Active to No (to retire a pillar, do that for all its pegs). Do not delete it. '
                         'Do not move a peg to another pillar mid-year: retire it and add a new one.').font = font(italic=True, color='667085')


    # ---------- Lists ----------
    ws = wb.create_sheet('Lists')
    blocks = [('A', 'Status', [s for s, _ in STATUSES], 'B', 'Meaning', [d for _, d in STATUSES], 18, 52),
              ('D', 'Format', list(formats), None, None, None, 14, 0),
              ('F', 'Team member', [t for t in team], None, None, None, 22, 0),
              ('H', 'Series', [s for s, _ in series], 'I', 'What it is', [d for _, d in series], 26, 48)]
    from openpyxl.utils import column_index_from_string as ci
    tnames = {'A': 'Statuses', 'D': 'Formats', 'F': 'Team', 'H': 'SeriesTable'}
    for col, head, vals, col2, head2, vals2, w1, w2 in blocks:
        c1 = ci(col)
        ws.cell(1, c1, head)
        ws.column_dimensions[col].width = w1
        if col2:
            ws.cell(1, ci(col2), head2)
            ws.column_dimensions[col2].width = w2
        style_header(ws, 1, c1, ci(col2) if col2 else c1)
        for r, v in enumerate(vals, start=2):
            ws.cell(r, c1, v).font = font()
            if col2:
                cc = ws.cell(r, ci(col2), vals2[r - 2])
                cc.font = font()
                cc.alignment = Alignment(wrap_text=True, vertical='top')
        end = max(2, len(vals) + 1)
        add_table(ws, tnames[col], f'{col}1:{col2 or col}{end}')
    for c in ('C', 'E', 'G'):
        ws.column_dimensions[c].width = 3
    ws.freeze_panes = 'A2'
    ws.cell(1, 11, 'Peg dropdown (auto)')
    for r in range(2, 302):
        ws.cell(r, 11, f'=IFERROR(INDEX(Pegs[Peg],MATCH(ROW()-1,Pegs[In dropdown (auto)],0)),"")')
    ws.column_dimensions['K'].hidden = True
    wb.defined_names['PegList'] = DefinedName('PegList', attr_text="OFFSET(Lists!$K$2,0,0,MAX(1,COUNTIF(Lists!$K$2:$K$301,\"?*\")),1)")
    list_name(wb, 'StatusList', 'Lists', 'A')
    list_name(wb, 'FormatList', 'Lists', 'D')
    list_name(wb, 'TeamList', 'Lists', 'F')
    list_name(wb, 'SeriesList', 'Lists', 'H')

    # ---------- Review flags (migration only) ----------
    if review or flags or peg_map:
        ws = wb.create_sheet('Review flags', 1)
        ws.sheet_view.showGridLines = False
        ws['A1'] = 'Review flags from the migration'
        ws['A1'].font = font(size=14, bold=True)
        ws['A2'] = ('Things to fix, grouped by person, then notes on what changed in the migration. Flagged cells are yellow; hover for the note. '
                    'Fix the cell, clear the yellow and mark it done here. Delete this sheet once everything is resolved.')
        ws['A2'].font = font(italic=True, color='667085')
        ws['A2'].alignment = Alignment(wrap_text=True)
        ws.merge_cells('A2:G2')
        ws.row_dimensions[2].height = 42
        heads = [('#', 5), ('For', 18), ('Sheet', 12), ('Cell', 8), ('Post', 30), ('What I found and what I did', 80), ('Done / decision', 26)]
        r0 = 4
        for i, (h, w) in enumerate(heads, start=1):
            ws.cell(r0, i, h)
            ws.column_dimensions[ws.cell(r0, i).column_letter].width = w
        style_header(ws, r0, 1, len(heads))
        items = [('Whole file', '', '', f'{f[0]} {f[1]}'.strip(), f[2] if len(f) > 2 else 'For information') for f in flags] + review
        items = sorted([x for x in items if x[4] != 'For information'], key=lambda x: (x[4] == 'BAU lead', x[4])) + \
            [x for x in items if x[4] == 'For information']
        for n, (sh, cell, title, text, who) in enumerate(items, start=1):
            r = r0 + n
            vals = [n, who, sh, cell, title, text, '']
            for i, v in enumerate(vals, start=1):
                c = ws.cell(r, i, v)
                c.font = font()
                c.alignment = Alignment(wrap_text=True, vertical='top')
                c.border = Border(bottom=THIN)
            if sh not in ('Whole file',) and cell:
                from openpyxl.worksheet.hyperlink import Hyperlink
                ws.cell(r, 4).hyperlink = Hyperlink(ref=f'D{r}', location=f"'{sh}'!{cell}")
                ws.cell(r, 4).font = font(color='1D4ED8', underline='single')
            if who != 'For information':
                ws.cell(r, 2).font = font(bold=True)
                ws.cell(r, 7).fill = PatternFill('solid', fgColor='FFFBEB')
        end = r0 + len(items)
        add_table(ws, 'ReviewFlags', f'A{r0}:G{end}')
        ws.freeze_panes = f'A{r0 + 1}'
        if peg_map:
            r = end + 3
            ws.cell(r, 1, 'How the old peg labels were mapped').font = font(size=12, bold=True)
            r += 1
            for i, h in enumerate(['#', 'Old sheet', 'Old pillar', 'Old peg label', 'New peg (pillar and comms task follow from it)', 'Posts'], start=1):
                ws.cell(r, i, h)
            style_header(ws, r, 1, 6)
            top = r
            for n, (old_pillar, old_peg, new_peg, count, sheets) in enumerate(peg_map, start=1):
                r += 1
                for i, v in enumerate([n, sheets, old_pillar, old_peg, new_peg, count], start=1):
                    c = ws.cell(r, i, v)
                    c.font = font()
                    c.alignment = Alignment(wrap_text=True, vertical='top')
                    c.border = Border(bottom=THIN)
            add_table(ws, 'PegMapping', f'A{top}:F{r}')

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return review


def options_from(path):
    """Reads the options (not the posts) from an existing workbook built by this script."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    start = wb['Start here']
    info = {start.cell(r, 1).value: start.cell(r, 2).value for r in range(3, 8)}

    def table_rows(sheet, name):
        ws = wb[sheet]
        t = ws.tables[name]
        rows = [[c.value for c in row] for row in ws[t.ref]]
        return [dict(zip(rows[0], r)) for r in rows[1:] if any(v not in (None, '') for v in r)]

    tasks = [(t['Code'], t['Comms task'], t['Counts as BAU'], t['Target share'], t['Active'], t['Notes'])
             for t in table_rows('Comms tasks', 'CommsTasks')]
    pegs = [(p['Comms task code'], p['Pillar'], p['Peg'], p['Active'], p['Notes'],
             wb['Framework'].cell(i + 2, 6).hyperlink.target if wb['Framework'].cell(i + 2, 6).hyperlink else p['Example'])
            for i, p in enumerate(table_rows('Framework', 'Pegs'))]
    bau = info.get('BAU started')
    return dict(
        brand=info['Brand'], purpose=info.get('Purpose') or '',
        bau_start=bau.date() if hasattr(bau, 'date') else bau,
        tasks=tasks, pegs=pegs,
        team=[t['Team member'] for t in table_rows('Lists', 'Team')],
        series=[(s['Series'], s['What it is']) for s in table_rows('Lists', 'SeriesTable')],
        formats=[f['Format'] for f in table_rows('Lists', 'Formats')],
    )


def main():
    import argparse
    ap = argparse.ArgumentParser(description='Build a blank content-calendar workbook for one brand and financial year.')
    ap.add_argument('--fy', type=int, required=True, help='Financial year by its ending year, e.g. 2027 for FY27 (April 2026 to March 2027).')
    ap.add_argument('--from', dest='source', help="Copy brand, framework and lists from an existing workbook (e.g. last year's).")
    ap.add_argument('--brand', help='Brand name, e.g. "Emkay Global". Required without --from.')
    ap.add_argument('--bau-start', help='Month the brand started the BAU framework, as YYYY-MM.')
    ap.add_argument('--out', help='Output file. Defaults to "<brand> - Content calendar FY<yy>.xlsx".')
    a = ap.parse_args()

    opts = options_from(a.source) if a.source else dict(tasks=[], pegs=[], team=[], series=[], formats=FORMATS)
    if a.brand:
        opts['brand'] = a.brand
    if a.bau_start:
        y, m = map(int, a.bau_start.split('-'))
        opts['bau_start'] = datetime.date(y, m, 1)
    if not opts.get('brand'):
        ap.error('--brand is required without --from')
    if not opts.get('tasks'):  # every brand tracks compliance posts outside the BAU mix
        opts['tasks'] = [('CO', 'Compliance', 'No', None, 'Yes', 'Mandatory or regulatory posts. Tracked, but not part of the BAU mix or targets.')]
        opts['pegs'] = [('CO', 'Mandatory communication', 'Investor awareness', 'Yes', 'Cybersecurity, fraud and safety alerts.', None),
                        ('CO', 'Mandatory communication', 'Regulatory disclosures', 'Yes', 'Notices the regulator requires us to publish.', None)]
    out = a.out or f"{opts['brand']} - Content calendar FY{a.fy % 100:02d}.xlsx"
    build(out, fy_start=a.fy - 1, **opts)
    print(f'Wrote {out}')


if __name__ == '__main__':
    main()
